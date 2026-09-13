"""Exporta la documentación PITA (Markdown) a Word (.docx) — estilo ejecutivo moderno.

Genera:
  - docs/DOCUMENTACION_MAESTRA_PITA.docx  (portada con banda lateral + TOC + cuerpo)
  - docs/DICCIONARIO_DE_DATOS.docx

Estilo: Calibri 11 justificado; títulos Calibri Light azul oscuro con línea inferior;
tablas "tipo libro" (solo líneas horizontales); código y notas con barra lateral azul.

Uso:  python docs/exportar_word.py
Requiere: pip install python-docx
Nota: al abrir el .docx, actualizar la tabla de contenido con clic derecho > "Actualizar campos".
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parent

AZUL = RGBColor(0x1F, 0x38, 0x64)        # azul institucional oscuro
AZUL_HEX = "1F3864"
GRIS = RGBColor(0x40, 0x40, 0x40)
GRIS_SUAVE_HEX = "F2F5FA"
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
FUENTE_CUERPO = "Calibri"
FUENTE_TITULO = "Calibri Light"
FUENTE_CODIGO = "Consolas"


# ---------------------------------------------------------------- utilidades
def _borde(elemento_pr, tag: str, estilo: str, tamano: int, color: str) -> None:
    borde = OxmlElement(f"w:{tag}")
    borde.set(qn("w:val"), estilo)
    borde.set(qn("w:sz"), str(tamano))
    borde.set(qn("w:space"), "0")
    borde.set(qn("w:color"), color)
    elemento_pr.append(borde)


def sombrear_celda(celda, color_hex: str) -> None:
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def sombrear_parrafo(parrafo, color_hex: str) -> None:
    p_pr = parrafo._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color_hex)
    p_pr.append(shd)


def barra_lateral(parrafo, color: str = AZUL_HEX, tamano: int = 18) -> None:
    """Borde izquierdo grueso: la 'barra' característica del estilo ejecutivo."""
    p_pr = parrafo._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    _borde(p_bdr, "left", "single", tamano, color)
    p_pr.append(p_bdr)


def linea_inferior(parrafo, color: str = AZUL_HEX, tamano: int = 8) -> None:
    p_pr = parrafo._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    _borde(p_bdr, "bottom", "single", tamano, color)
    p_pr.append(p_bdr)


def tabla_estilo_libro(tabla) -> None:
    """Estilo moderno: cabeceras marcadas y líneas tenues."""
    tbl_pr = tabla._tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    _borde(bordes, "top", "single", 12, AZUL_HEX)
    _borde(bordes, "bottom", "single", 12, AZUL_HEX)
    _borde(bordes, "insideH", "single", 4, "E0E0E0")  # Gris más claro y sutil
    for lado in ("left", "right", "insideV"):
        _borde(bordes, lado, "none", 0, "auto")
    tbl_pr.append(bordes)


def sin_bordes(tabla) -> None:
    tbl_pr = tabla._tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "bottom", "left", "right", "insideH", "insideV"):
        _borde(bordes, lado, "none", 0, "auto")
    tbl_pr.append(bordes)


def agregar_campo_toc(doc: Document) -> None:
    parrafo = doc.add_paragraph()
    run = parrafo.add_run()
    inicio = OxmlElement("w:fldChar"); inicio.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fin = OxmlElement("w:fldChar"); fin.set(qn("w:fldCharType"), "end")
    run._r.append(inicio); run._r.append(instr); run._r.append(fin)


def numero_pagina_footer(doc: Document) -> None:
    parrafo = doc.sections[0].footer.paragraphs[0]
    parrafo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = parrafo.add_run("Página ")
    run.font.size = Pt(9)
    run.font.color.rgb = GRIS
    run2 = parrafo.add_run()
    inicio = OxmlElement("w:fldChar"); inicio.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fin = OxmlElement("w:fldChar"); fin.set(qn("w:fldCharType"), "end")
    run2._r.append(inicio); run2._r.append(instr); run2._r.append(fin)
    run2.font.size = Pt(9)
    run2.font.color.rgb = GRIS


def configurar_estilos(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FUENTE_CUERPO
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x21, 0x21, 0x21)
    normal.paragraph_format.space_after = Pt(10)  # Aumentado para que el texto respire
    normal.paragraph_format.line_spacing = 1.15

    tamanos = {1: 22, 2: 16, 3: 13, 4: 11.5}
    for nivel, tamano in tamanos.items():
        estilo = doc.styles[f"Heading {nivel}"]
        estilo.font.name = FUENTE_TITULO
        estilo.font.size = Pt(tamano)
        estilo.font.bold = nivel >= 2
        estilo.font.color.rgb = AZUL
        
        if nivel == 1:
            estilo.paragraph_format.space_before = Pt(28)
            estilo.paragraph_format.space_after = Pt(12)
        elif nivel == 2:
            estilo.paragraph_format.space_before = Pt(22)
            estilo.paragraph_format.space_after = Pt(8)
        else:
            estilo.paragraph_format.space_before = Pt(16)
            estilo.paragraph_format.space_after = Pt(6)
            
        estilo.paragraph_format.keep_with_next = True


def encabezado(doc: Document, texto: str, nivel: int):
    p = doc.add_heading(texto, level=nivel)
    if nivel == 1:
        linea_inferior(p, tamano=12)
    elif nivel == 2:
        linea_inferior(p, tamano=6)
    return p


# ------------------------------------------------------------- formato inline
PATRON_INLINE = re.compile(r"(\*\*.+?\*\*|~~.+?~~|`.+?`)")


def agregar_runs(parrafo, texto: str) -> None:
    for parte in PATRON_INLINE.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            run = parrafo.add_run(parte[2:-2]); run.bold = True
        elif parte.startswith("~~") and parte.endswith("~~"):
            run = parrafo.add_run(parte[2:-2]); run.font.strike = True
        elif parte.startswith("`") and parte.endswith("`"):
            run = parrafo.add_run(parte[1:-1])
            run.font.name = FUENTE_CODIGO
            run.font.size = Pt(9.5)
            run.font.color.rgb = AZUL
        else:
            parrafo.add_run(parte)


# ------------------------------------------------------------------ portadas
def _contenido_portada(celda, lineas_gris: list[str], titulo: str, subtitulo: str | None) -> None:
    """Rellena la celda derecha de la portada (a la izquierda va la banda azul)."""
    primera = celda.paragraphs[0]
    primera.paragraph_format.space_before = Pt(160)
    for i, linea in enumerate(lineas_gris):
        p = primera if i == 0 else celda.add_paragraph()
        run = p.add_run(linea)
        run.font.name = FUENTE_TITULO
        run.font.size = Pt(12)
        run.font.color.rgb = GRIS
        p.paragraph_format.space_after = Pt(2)
    p = celda.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    run = p.add_run(titulo.upper())
    run.font.name = FUENTE_TITULO
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = AZUL
    linea_inferior(p, tamano=16)
    if subtitulo:
        p = celda.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        run = p.add_run(subtitulo)
        run.font.size = Pt(12)
        run.font.color.rgb = GRIS
    p = celda.add_paragraph()
    p.paragraph_format.space_before = Pt(120)
    run = p.add_run("Septiembre de 2026 — Valledupar, Cesar")
    run.font.size = Pt(11)
    run.font.color.rgb = GRIS


def portada(doc: Document, lineas_gris: list[str], titulo: str, subtitulo: str | None = None) -> None:
    """Portada con banda lateral azul: tabla invisible de 2 columnas."""
    tabla = doc.add_table(rows=1, cols=2)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    sin_bordes(tabla)
    banda, contenido = tabla.rows[0].cells
    banda.width = Cm(0.9)
    contenido.width = Cm(15.1)
    sombrear_celda(banda, AZUL_HEX)
    banda.paragraphs[0].add_run("\n" * 26)  # altura de la banda
    _contenido_portada(contenido, lineas_gris, titulo, subtitulo)
    doc.add_page_break()


def datos_portada_maestra(lineas_md: list[str]) -> tuple[list[str], str]:
    datos = []
    iniciado = False
    for l in lineas_md:
        texto = l.strip()
        if texto.startswith(">"):
            iniciado = True
            limpio = texto.lstrip(">").strip().rstrip("  ")
            limpio = re.sub(r"\*\*", "", limpio)
            if limpio:
                datos.append(limpio)
        elif iniciado and not texto.startswith(">"):
            break
    titulo = lineas_md[0].lstrip("# ").strip()
    return datos, titulo


INSTITUCION = [
    "Universidad Popular del Cesar",
    "Facultad de Ingenierías y Tecnológicas",
    "Ingeniería de Sistemas — Estructura de Datos",
    "Taller 1, Listas — Sistema PITA",
]


# ------------------------------------------------------------------ parser
def convertir(md_path: Path, docx_path: Path, portada_tipo: str = "maestra") -> None:
    lineas = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configurar_estilos(doc)
    for seccion in doc.sections:
        seccion.top_margin = seccion.bottom_margin = Cm(2.5)
        seccion.left_margin = seccion.right_margin = Cm(2.5)
    numero_pagina_footer(doc)

    if portada_tipo == "maestra":
        datos, titulo = datos_portada_maestra(lineas)
        portada(doc, datos, titulo)
    else:
        portada(doc, INSTITUCION, "Diccionario de datos del sistema PITA",
                "Anexo del Documento Maestro PITA — 30 entidades y 607 atributos")

    encabezado(doc, "Tabla de contenido", 1)
    agregar_campo_toc(doc)
    doc.add_page_break()

    # Omitir del cuerpo el bloque de encabezado ya usado en la portada
    inicio_cuerpo = 0
    # Omitir el título principal
    while inicio_cuerpo < len(lineas):
        if lineas[inicio_cuerpo].startswith("# "):
            inicio_cuerpo += 1
            break
        inicio_cuerpo += 1
    
    # Omitir bloque de citas y separadores iniciales
    while inicio_cuerpo < len(lineas) and (
        lineas[inicio_cuerpo].strip().startswith(">")
        or not lineas[inicio_cuerpo].strip() 
        or lineas[inicio_cuerpo].strip() == "---"
    ):
        inicio_cuerpo += 1

    i = inicio_cuerpo
    while i < len(lineas):
        linea = lineas[i]
        texto = linea.strip()

        if not texto:
            i += 1
            continue

        if texto.startswith("```"):
            bloque = []
            i += 1
            while i < len(lineas) and not lineas[i].strip().startswith("```"):
                bloque.append(lineas[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            sombrear_parrafo(p, GRIS_SUAVE_HEX)
            barra_lateral(p)
            p.paragraph_format.left_indent = Cm(0.3)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(10)
            run = p.add_run("\n".join(bloque))
            run.font.name = FUENTE_CODIGO
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(0x2A, 0x2A, 0x2A)
            continue

        if texto.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                celdas = [c.strip() for c in lineas[i].strip().strip("|").split("|")]
                filas.append(celdas)
                i += 1
            filas = [f for f in filas if not all(re.fullmatch(r":?-+:?", c or "-") for c in f)]
            if filas:
                tabla = doc.add_table(rows=len(filas), cols=len(filas[0]))
                tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
                tabla_estilo_libro(tabla)
                for f, fila in enumerate(filas):
                    es_cabecera = (f == 0)
                    es_fila_par = (f % 2 == 0) and not es_cabecera
                    
                    for c, celda_texto in enumerate(fila):
                        if c >= len(tabla.rows[f].cells):
                            continue
                        celda = tabla.rows[f].cells[c]
                        
                        if es_cabecera:
                            sombrear_celda(celda, AZUL_HEX)
                        elif es_fila_par:
                            sombrear_celda(celda, "F9FAFB") # Efecto cebra muy sutil
                            
                        p = celda.paragraphs[0]
                        p.text = ""
                        p.paragraph_format.space_before = Pt(4)
                        p.paragraph_format.space_after = Pt(4)
                        
                        agregar_runs(p, celda_texto)
                        for run in p.runs:
                            run.font.size = Pt(9.5)
                            if es_cabecera:
                                run.font.bold = True
                                run.font.color.rgb = BLANCO
                            else:
                                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                doc.add_paragraph()
            continue

        encabezado_md = re.match(r"^(#{1,6})\s+(.*)$", texto)
        if encabezado_md:
            nivel = min(len(encabezado_md.group(1)), 4)
            p_enc = encabezado(doc, re.sub(r"\*\*", "", encabezado_md.group(2)), nivel)
            if nivel == 1:
                p_enc.paragraph_format.page_break_before = True
            i += 1
            continue

        if texto.startswith(">"):
            nota = []
            while i < len(lineas) and lineas[i].strip().startswith(">"):
                nota.append(lineas[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            barra_lateral(p)
            p.paragraph_format.left_indent = Cm(0.4)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(10)
            agregar_runs(p, " ".join(nota))
            for run in p.runs:
                run.font.size = Pt(10)
                run.font.color.rgb = GRIS
            continue

        if texto in ("---", "***"):
            i += 1
            continue

        viñeta = re.match(r"^[-*]\s+(.*)$", texto)
        numerada = re.match(r"^\d+\.\s+(.*)$", texto)
        if viñeta:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(6)
            agregar_runs(p, viñeta.group(1))
            i += 1
            continue
        if numerada:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(6)
            agregar_runs(p, numerada.group(1))
            i += 1
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        agregar_runs(p, texto)
        i += 1

    doc.save(docx_path)
    print(f"Generado: {docx_path}")


if __name__ == "__main__":
    # Se exportan las versiones limpias generadas por limpiar_md.py
    convertir(DOCS / "export" / "DOCUMENTACION_MAESTRA_PITA.md", DOCS / "DOCUMENTACION_MAESTRA_PITA.docx", "maestra")
    convertir(DOCS / "export" / "DICCIONARIO_DE_DATOS.md", DOCS / "DICCIONARIO_DE_DATOS.docx", "anexo")
